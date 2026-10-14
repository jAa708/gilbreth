"""Salidas del agente: filas del diagrama (uniendo segundos iguales), indicadores, Excel en el formato del curso
(cuatro columnas de símbolos por mano, resumen con fórmulas) y una página HTML de revisión con un fotograma por fila.
Las propuestas de mejora salen de reglas de economía de movimientos aplicadas a los números, no de texto libre."""
import base64, html, json, re
from collections import Counter
from pathlib import Path

SIMBOLO = {'OP': 'O', 'TR+OP': 'O', 'TR': 'T', 'S': 'S', 'D': 'E', 'NV': 'E'}
NOMBRE = {'OP': 'operación', 'TR': 'transporte', 'TR+OP': 'transporte y operación', 'S': 'sostiene', 'D': 'espera',
          'NV': 'no visible'}
NOMBRE_EN = {'OP': 'Operation', 'TR': 'Transport', 'TR+OP': 'Transport + operation', 'S': 'Hold', 'D': 'Delay',
             'NV': 'Not visible'}


def nombre(cod, idioma='es'):
    return (NOMBRE_EN if idioma == 'en' else NOMBRE)[cod]


# English wording for metadata in the two older, Spanish-only example files. New
# charts already supply their metadata in the language selected by the agent.
_META_EN = {
    'Caja reductora (IMPACT)': 'Gearbox (IMPACT)',
    'Rearmado de una caja reductora con herramientas': 'Gearbox reassembly with tools',
    'Puesto de taller, cámara cenital': 'Workshop bench, overhead camera',
    'LEGO Creator 3-en-1 nº 5761 «Mini Digger», modelo 1 (56 piezas)':
        'LEGO Creator 3-in-1 No. 5761 Mini Digger, model 1 (56 pieces)',
    'Armado completo siguiendo el folleto. v7: calibrado con el minuto de oro revisado (0-70 s)':
        'Full assembly following the instruction booklet. v7: calibrated on the reviewed first 70 s',
    'Mesa con piezas clasificadas por color; folleto de instrucciones':
        'Table with parts sorted by color; instruction booklet',
    'Convención del curso: ▽ solo si la otra mano opera.':
        'Course convention: ▽ only when the other hand operates.',
    '▽ solo si la otra mano opera; sostener mientras la otra va por la pieza o folleto = D (curso).':
        '▽ only when the other hand operates; holding while the other hand fetches a part or the booklet = D (course).',
}


def _meta_text(meta, clave, idioma='es'):
    valor = meta.get(f'{clave}_en', meta.get(clave, '')) if idioma == 'en' else meta.get(clave, '')
    return _META_EN.get(valor, valor) if idioma == 'en' else valor


def _descripcion(f, clave, idioma='es'):
    valor = f.get(f'desc_{clave}') or nombre(f[clave], idioma)
    if idioma == 'en' and valor in NOMBRE.values():
        return NOMBRE_EN[next(cod for cod, texto in NOMBRE.items() if texto == valor)]
    return valor


def mmss(t):
    return f'{int(t // 60)}:{t % 60:04.1f}'.replace('.0', '') if t % 1 else f'{int(t // 60)}:{int(t % 60):02d}'


def filas_elemento(ventanas, max_tr=1):
    """Nivel «elemento» (Mundel): si una mano transporta ≤ max_tr s e inmediatamente opera, ese transporte se une a la
    operación como TR+OP («trae y encaja»), y la otra mano que sostenía pasa de D a S (convención del curso). La OIT
    estricta (por defecto) deja transporte y operación en filas distintas salvo simultaneidad real."""
    v = [dict(o) for o in sorted(ventanas, key=lambda x: x['t0'])]
    for k, otra in (('md', 'mi'), ('mi', 'md')):
        i = 0
        while i < len(v):
            if v[i][k] == 'TR':
                j = i
                while j + 1 < len(v) and v[j + 1][k] == 'TR':
                    j += 1
                if j - i + 1 <= max_tr and j + 1 < len(v) and v[j + 1][k] in ('OP', 'TR+OP'):
                    for t in range(i, j + 1):
                        v[t][k] = 'TR+OP'
                        if v[t][otra] == 'D':
                            v[t][otra] = 'S'
                i = j + 1
            else:
                i += 1
    return v


def unir_filas(ventanas):
    """ventanas: [{'t0','t1','md','mi','crudo'}] en orden → filas del diagrama (una fila por cambio de cualquier mano)."""
    filas = []
    for v in sorted(ventanas, key=lambda x: x['t0']):
        obs = (v.get('crudo') or {}).get('observations', '')
        if filas and filas[-1]['md'] == v['md'] and filas[-1]['mi'] == v['mi'] and abs(filas[-1]['t1'] - v['t0']) < 0.01:
            filas[-1]['t1'] = v['t1']
            filas[-1]['obs'].append(obs)
            filas[-1]['ks'].append(v.get('k'))
        else:
            filas.append({'t0': v['t0'], 't1': v['t1'], 'md': v['md'], 'mi': v['mi'], 'obs': [obs], 'ks': [v.get('k')]})
    for i, f in enumerate(filas, 1):
        f['n'] = i
        f['dur'] = round(f['t1'] - f['t0'], 2)
    return filas


def indicadores(filas):
    T = sum(f['dur'] for f in filas) or 1
    out = {'duracion_s': round(T, 1), 'filas': len(filas)}
    for mano, k in (('MD', 'md'), ('MI', 'mi')):
        c = Counter()
        for f in filas:
            c[f[k]] += f['dur']
        out[mano] = {cod: round(100 * s / T, 1) for cod, s in c.most_common()}
        out[mano + '_productivo_pct'] = round(100 * (c['OP'] + c['TR+OP']) / T, 1)
        out[mano + '_improductivo_pct'] = round(100 * (c['D'] + c['S'] + c['NV']) / T, 1)
    return out


def propuestas(ind, idioma='es'):
    """Reglas de economía de movimientos (OIT / Barnes) disparadas por los números del diagrama."""
    p = []
    for mano, otra in (('MI', 'MD'), ('MD', 'MI')):
        s, d = ind[mano].get('S', 0), ind[mano].get('D', 0)
        if idioma == 'en':
            hand = 'left hand' if mano == 'MI' else 'right hand'
            if s >= 25:
                p.append(f'The {hand} holds the workpiece for {s:.0f}% of the observed time. Use a fixture or jig '
                         f'to hold it and free the hand for simultaneous work; a hand should not serve as a clamp.')
            if d >= 20:
                p.append(f'The {hand} is delayed for {d:.0f}% of the observed time. Balance the work so both hands '
                         f'start and finish together and move symmetrically.')
            tr = ind[mano].get('TR', 0)
            if tr >= 20:
                p.append(f'The {hand} transports for {tr:.0f}% of the observed time. Place parts and tools within '
                         f'normal reach, in order of use, and consider gravity feed.')
            continue
        if s >= 25:
            p.append(f'La {"izquierda" if mano == "MI" else "derecha"} sostiene el {s:.0f} % del tiempo: un soporte o '
                     f'plantilla que sujete la pieza la libera para trabajar en simultáneo (principio: no usar la mano como '
                     f'mordaza).')
        if d >= 20:
            p.append(f'La {"izquierda" if mano == "MI" else "derecha"} espera el {d:.0f} % del tiempo: repartir el trabajo '
                     f'para que ambas manos empiecen y terminen a la vez y hagan movimientos simétricos.')
        tr = ind[mano].get('TR', 0)
        if tr >= 20:
            p.append(f'La {"izquierda" if mano == "MI" else "derecha"} transporta el {tr:.0f} % del tiempo: acercar '
                     f'piezas y herramientas al área normal de trabajo, en el orden de uso y con alimentación por gravedad.')
    if abs(ind['MD_productivo_pct'] - ind['MI_productivo_pct']) >= 30:
        if idioma == 'en':
            p.append(f'Uneven workload: the right hand operates for {ind["MD_productivo_pct"]:.0f}% of the observed '
                     f'time and the left for {ind["MI_productivo_pct"]:.0f}%. Check whether the left hand can feed '
                     f'parts while the right hand assembles.')
        else:
            p.append(f'Desbalance: la derecha opera el {ind["MD_productivo_pct"]:.0f} % y la izquierda el '
                     f'{ind["MI_productivo_pct"]:.0f} %. Revisar si la izquierda puede alimentar piezas mientras la derecha ensambla.')
    return p or (['No major imbalance detected. Review delays and long transports row by row.'] if idioma == 'en' else
                 ['Sin desbalances grandes: revisar fila por fila las esperas y transportes largos.'])


def _detalle_error(e, idioma='es'):
    if not e.get('resumen'):
        return f'{mmss(e["t0"])} – {mmss(e["t1"])}'
    if idioma == 'es':
        return e['detalle']
    m = re.match(r'^(\d+) veces, ([\d.]+) s en total', e['detalle'])
    return (f'{m.group(1)} occurrences, {m.group(2)} s total (throughout the cycle)' if m else
            f'{mmss(e["t0"])} – {mmss(e["t1"])}')


def escribir_excel(filas, ind, meta, ruta, eventos=None, idioma='es'):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = Workbook()
    ws = wb.active
    ws.title = 'Bimanual chart' if idioma == 'en' else 'Diagrama bimanual'
    AZUL, GRIS, AMAR = PatternFill('solid', fgColor='DBE5F1'), PatternFill('solid', fgColor='F2F2F2'), PatternFill('solid', fgColor='FFF2CC')
    fino = Side(style='thin', color='808080')
    BORDE = Border(left=fino, right=fino, top=fino, bottom=fino)
    F = lambda b=False, s=10, i=False: Font(name='Arial', size=s, bold=b, italic=i)
    C = Alignment(horizontal='center', vertical='center', wrap_text=True)
    L = Alignment(horizontal='left', vertical='center', wrap_text=True)
    for k, v in {'A': 5, 'B': 30, 'C': 22, 'D': 8, 'E': 5, 'F': 5, 'G': 5, 'H': 5, 'I': 2, 'J': 5, 'K': 5, 'L': 5, 'M': 5,
                 'N': 8, 'O': 30, 'P': 22, 'Q': 5}.items():
        ws.column_dimensions[k].width = v
    ws.merge_cells('A1:Q1')
    ws['A1'] = 'B I M A N U A L   C H A R T' if idioma == 'en' else 'D I A G R A M A   B I M A N U A L'
    ws['A1'].font, ws['A1'].alignment, ws['A1'].fill = F(True, 14), C, AZUL
    if idioma == 'en':
        cab = [('Chart No. 1 · Sheet 1 of 1', 'Method: CURRENT (observed on video)'),
               (f'Product: {_meta_text(meta, "producto", idioma)}'[:250],
                f'Video: {meta.get("video", "")} ({ind["duracion_s"]} s)'),
               (f'Operation: {_meta_text(meta, "operacion", idioma)}'[:250],
                f'Workstation: {_meta_text(meta, "lugar", idioma)}'[:250]),
               ('Prepared by: bimanual chart agent (draft for analyst review)', f'Date: {meta.get("fecha", "")}')]
    else:
        cab = [('Diagrama núm. 1 · Hoja 1 de 1', 'Método: ACTUAL (observado en video)'),
               (f'Producto: {meta.get("producto", "")}'[:250], f'Video: {meta.get("video", "")} ({ind["duracion_s"]} s)'),
               (f'Operación: {meta.get("operacion", "")}'[:250], f'Lugar: {meta.get("lugar", "")}'[:250]),
               ('Elaborado: agente bimanual (borrador para revisión del analista)', f'Fecha: {meta.get("fecha", "")}')]
    for i, (a, b) in enumerate(cab, start=2):
        ws.merge_cells(f'A{i}:H{i}'), ws.merge_cells(f'I{i}:Q{i}')
        ws[f'A{i}'], ws[f'I{i}'] = a, b
        for c in (ws[f'A{i}'], ws[f'I{i}']):
            c.font, c.alignment = F(s=9), L
        ws.row_dimensions[i].height = 30
    h = 7
    ws.merge_cells(f'B{h}:C{h}'), ws.merge_cells(f'O{h}:P{h}')
    if idioma == 'en':
        ws[f'A{h}'], ws[f'B{h}'], ws[f'D{h}'] = 'No.', 'Right-hand description (R)', 'Time (s)'
        ws[f'N{h}'], ws[f'O{h}'], ws[f'Q{h}'] = 'Time (s)', 'Left-hand description (L)', 'No.'
    else:
        ws[f'A{h}'], ws[f'B{h}'], ws[f'D{h}'] = 'N.º', 'Descripción de la mano derecha (MD)', 'Tiempo (s)'
        ws[f'N{h}'], ws[f'O{h}'], ws[f'Q{h}'] = 'Tiempo (s)', 'Descripción de la mano izquierda (MI)', 'N.º'
    for cols in ('EFGH', 'JKLM'):
        for col, sym in zip(cols, ['○', '⇨', 'D', '▽']):
            ws[f'{col}{h}'] = sym
    for col in 'ABCDEFGHJKLMNOPQ':
        c = ws[f'{col}{h}']
        c.font, c.alignment, c.fill, c.border = F(True, 9), C, AZUL, BORDE
    r0 = h + 1
    for k, f in enumerate(filas):
        r = r0 + k
        ws.merge_cells(f'B{r}:C{r}'), ws.merge_cells(f'O{r}:P{r}')
        ws[f'A{r}'] = ws[f'Q{r}'] = f['n']
        ws[f'B{r}'] = f'{mmss(f["t0"])}–{mmss(f["t1"])}  {_descripcion(f, "md", idioma)}'
        ws[f'O{r}'] = f'{mmss(f["t0"])}–{mmss(f["t1"])}  {_descripcion(f, "mi", idioma)}'
        ws[f'D{r}'] = ws[f'N{r}'] = f['dur']
        ws[f'{dict(O="E", T="F", E="G", S="H")[SIMBOLO[f["md"]]]}{r}'] = 'x'
        ws[f'{dict(O="J", T="K", E="L", S="M")[SIMBOLO[f["mi"]]]}{r}'] = 'x'
        for col in 'ABCDEFGHJKLMNOPQ':
            c = ws[f'{col}{r}']
            c.border, c.font = BORDE, F(s=9)
            c.alignment = L if col in 'BCOP' else C
            if c.value == 'x':
                c.fill, c.font = AMAR, F(True, 10)
        ws.row_dimensions[r].height = 28
    rf = r0 + len(filas) - 1
    rs = rf + 2
    ws.merge_cells(f'A{rs}:Q{rs}')
    ws[f'A{rs}'] = 'SUMMARY' if idioma == 'en' else 'RESUMEN'
    ws[f'A{rs}'].font, ws[f'A{rs}'].alignment, ws[f'A{rs}'].fill = F(True, 11), C, AZUL
    if idioma == 'en':
        ws[f'B{rs + 1}'], ws[f'D{rs + 1}'], ws[f'E{rs + 1}'], ws[f'J{rs + 1}'], ws[f'K{rs + 1}'] = (
            'Current method', 'R (count)', 'R (s)', 'L (count)', 'L (s)')
        resumen = [('Operations ○', 'E', 'J'), ('Transports ⇨', 'F', 'K'),
                   ('Delays D', 'G', 'L'), ('Holds ▽', 'H', 'M')]
    else:
        ws[f'B{rs + 1}'], ws[f'D{rs + 1}'], ws[f'E{rs + 1}'], ws[f'J{rs + 1}'], ws[f'K{rs + 1}'] = 'Método actual', 'MD (n.º)', 'MD (s)', 'MI (n.º)', 'MI (s)'
        resumen = [('Operaciones ○', 'E', 'J'), ('Desplazamientos ⇨', 'F', 'K'),
                   ('Esperas D', 'G', 'L'), ('Sostenimientos ▽', 'H', 'M')]
    for i, (nom, cd, ci) in enumerate(resumen, start=rs + 2):
        ws[f'B{i}'] = nom
        ws[f'D{i}'] = f'=COUNTIF({cd}{r0}:{cd}{rf},"x")'
        ws[f'E{i}'] = f'=SUMIF({cd}{r0}:{cd}{rf},"x",$D${r0}:$D${rf})'
        ws[f'J{i}'] = f'=COUNTIF({ci}{r0}:{ci}{rf},"x")'
        ws[f'K{i}'] = f'=SUMIF({ci}{r0}:{ci}{rf},"x",$N${r0}:$N${rf})'
        for c in 'BDEJK':
            ws[f'{c}{i}'].font, ws[f'{c}{i}'].border = F(s=9), BORDE
    rp = rs + 7
    ws.merge_cells(f'A{rp}:Q{rp}')
    ws[f'A{rp}'] = ('IMPROVEMENT PROPOSALS (motion-economy rules applied to the chart data)' if idioma == 'en' else
                   'PROPUESTAS DE MEJORA (reglas de economía de movimientos sobre los números del diagrama)')
    ws[f'A{rp}'].font, ws[f'A{rp}'].alignment, ws[f'A{rp}'].fill = F(True, 11), C, AZUL
    prop = propuestas(ind, idioma)
    for i, t in enumerate(prop, start=rp + 1):
        ws.merge_cells(f'A{i}:Q{i}')
        ws[f'A{i}'] = f'{i - rp}) {t}'
        ws[f'A{i}'].font, ws[f'A{i}'].alignment = F(s=9), L
        ws.row_dimensions[i].height = 28
    nota = rp + len(prop) + 2
    ws.merge_cells(f'A{nota}:Q{nota}')
    if idioma == 'en':
        ws[f'A{nota}'] = (f'Convention: {_meta_text(meta, "convencion_texto", idioma)} '
                          f'Automatic draft: the analyst reviews and approves each row before use. '
                          f'Observed time (not normal or standard time: performance rating and allowances have not been applied).')
    else:
        ws[f'A{nota}'] = (f'Convención: {meta.get("convencion_texto", "")} Borrador automático: el analista revisa y aprueba cada fila '
                          f'antes de usarlo. Tiempo observado (no normal ni estándar: falta la valoración del ritmo y los suplementos).')
    ws[f'A{nota}'].font, ws[f'A{nota}'].alignment = F(s=8, i=True), L
    ws.row_dimensions[nota].height = 36
    ws.freeze_panes = f'A{r0}'
    ws.page_setup.orientation = 'landscape'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    if meta.get('licencia_en'):
        wl = wb.create_sheet('Source and licence' if idioma == 'en' else 'Fuente y licencia')
        wl.column_dimensions['A'].width = 40
        for col in 'BCDEFGH':
            wl.column_dimensions[col].width = 10
        wl.merge_cells('A1:H1')
        wl['A1'] = 'Source and licence' if idioma == 'en' else 'Fuente y licencia'
        wl['A1'].font, wl['A1'].alignment, wl['A1'].fill = F(True, 14), L, AZUL
        wl.merge_cells('A2:H5')
        wl['A2'] = meta['licencia_en']
        wl['A2'].font, wl['A2'].alignment = F(s=11), Alignment(horizontal='left', vertical='top', wrap_text=True)
        for row in range(2, 6):
            wl.row_dimensions[row].height = 30
        wl.sheet_properties.pageSetUpPr.fitToPage = True
        wl.page_setup.fitToWidth, wl.page_setup.fitToHeight = 1, 1
    if eventos:
        import errores
        we = wb.create_sheet('Method errors' if idioma == 'en' else 'Errores de método')
        for k, v in {'A': 5, 'B': 34, 'C': 12, 'D': 14, 'E': 10, 'F': 60}.items():
            we.column_dimensions[k].width = v
        we.merge_cells('A1:F1')
        we['A1'] = ('POSSIBLE METHOD ERRORS (cues to check against the video, not final findings)' if idioma == 'en' else
                    'POSIBLES ERRORES DE MÉTODO (señales para revisar contra el video; no son veredictos)')
        we['A1'].font, we['A1'].alignment, we['A1'].fill = F(True, 12), C, AZUL
        encabezados = (['No.', 'Error', 'Hand', 'From – to', 'Duration (s)', 'Suggested correction'] if idioma == 'en' else
                       ['N.º', 'Error', 'Mano', 'Desde – hasta', 'Duración (s)', 'Corrección propuesta'])
        for col, h in zip('ABCDEF', encabezados):
            we[f'{col}2'] = h
            we[f'{col}2'].font, we[f'{col}2'].alignment, we[f'{col}2'].fill, we[f'{col}2'].border = F(True, 9), C, GRIS, BORDE
        for i, e in enumerate(errores.para_informe(eventos), 1):
            r = i + 2
            manos = ({'R': 'right', 'L': 'left', 'RL': 'both', 'LR': 'both'} if idioma == 'en' else
                     {'R': 'derecha', 'L': 'izquierda', 'RL': 'ambas', 'LR': 'ambas'})
            mano = manos.get(e['mano'], e['mano'])
            cuando = _detalle_error(e, idioma)
            dur = float(e['detalle'].split(' veces, ')[1].split(' s')[0]) if e.get('resumen') else round(e['t1'] - e['t0'], 1)
            titulo = errores.TITLE_EN if idioma == 'en' else errores.TITULO
            correccion = errores.FIX_EN if idioma == 'en' else errores.CORRECCION
            vals = [i, titulo.get(e['tipo'], e['tipo']), mano, cuando, dur, correccion.get(e['tipo'], '')]
            for col, v in zip('ABCDEF', vals):
                c = we[f'{col}{r}']
                c.value, c.font, c.border = v, F(s=9), BORDE
                c.alignment = L if col in 'BF' else C
            we.row_dimensions[r].height = 30
    wb.save(ruta)


def escribir_html(filas, ind, meta, carpeta_img, ruta, img_de=None, eventos=None, idioma='es'):
    """Página de revisión: una tarjeta por fila con el fotograma central, los códigos y lo que vio el modelo.
    img_de(k) -> ruta del fotograma (por defecto carpeta_img/{k:04d}_f1.jpg)."""
    def img(k):
        p = Path(img_de(k)) if img_de else Path(carpeta_img) / f'{k:04d}_f1.jpg'
        return 'data:image/jpeg;base64,' + base64.b64encode(p.read_bytes()).decode() if p.exists() else ''
    tarjetas = []
    for f in filas:
        k = f['ks'][len(f['ks']) // 2]
        obs = ' · '.join(dict.fromkeys(o for o in f['obs'] if o))[:500]
        if idioma == 'en':
            detalle = (f'R <span class="c">{f["md"]}</span> ({nombre(f["md"], idioma)}) · '
                       f'L <span class="c">{f["mi"]}</span> ({nombre(f["mi"], idioma)})<br>'
                       f'<small>R: {html.escape(_descripcion(f, "md", idioma))} · '
                       f'L: {html.escape(_descripcion(f, "mi", idioma))}</small>')
        else:
            detalle = (f'MD <span class="c">{f["md"]}</span> · MI <span class="c">{f["mi"]}</span><br>'
                       f'<small>{html.escape(obs)}</small>')
        tarjetas.append(f'''<div class="t"><img src="{img(k)}" loading="lazy"><div class="d">
<b>{f["n"]}. {mmss(f["t0"])}–{mmss(f["t1"])} ({f["dur"]} s)</b><br>
{detalle}</div></div>''')
    pct = (lambda m: ' · '.join(f'{k} ({nombre(k, idioma)}) {v:.0f} %' for k, v in ind[m].items())) if idioma == 'en' else (
        lambda m: ' · '.join(f'{k} {v:.0f} %' for k, v in ind[m].items()))
    prop = ''.join(f'<li>{html.escape(t)}</li>' for t in propuestas(ind, idioma))
    err = ''
    if eventos:
        import errores
        manos = ({'R': 'right', 'L': 'left', 'RL': 'both', 'LR': 'both'} if idioma == 'en' else
                 {'R': 'derecha', 'L': 'izquierda', 'RL': 'ambas', 'LR': 'ambas'})
        titulos = errores.TITLE_EN if idioma == 'en' else errores.TITULO
        correcciones = errores.FIX_EN if idioma == 'en' else errores.CORRECCION
        filas_err = ''.join(f'<tr><td>{html.escape(titulos.get(e["tipo"], e["tipo"]))}</td><td>{manos.get(e["mano"], e["mano"])}</td>'
                            f'<td>{html.escape(_detalle_error(e, idioma)) if idioma == "en" else html.escape(e["detalle"]) if e.get("resumen") else mmss(e["t0"]) + "–" + mmss(e["t1"])}</td>'
                            f'<td>{html.escape(correcciones.get(e["tipo"], ""))}</td></tr>'
                            for e in errores.para_informe(eventos))
        if idioma == 'en':
            err = (f'<h2>Possible method errors (cues to review)</h2><table><tr><th>Error</th><th>Hand</th>'
                   f'<th>When</th><th>Suggested correction</th></tr>{filas_err}</table>')
        else:
            err = (f'<h2>Posibles errores de método (señales para revisar)</h2><table><tr><th>Error</th><th>Mano</th>'
                   f'<th>Cuándo</th><th>Corrección propuesta</th></tr>{filas_err}</table>')
    idioma_html = 'en' if idioma == 'en' else 'es'
    titulo = 'Bimanual chart review' if idioma == 'en' else 'Revisión del diagrama bimanual'
    encabezado = 'Bimanual chart: review' if idioma == 'en' else 'Diagrama bimanual: revisión'
    n_filas = 'rows' if idioma == 'en' else 'filas'
    mano_d, mano_i = ('R', 'L') if idioma == 'en' else ('MD', 'MI')
    propuestas_titulo = 'Improvement proposals' if idioma == 'en' else 'Propuestas'
    filas_titulo = 'Rows' if idioma == 'en' else 'Filas'
    ruta = Path(ruta)
    ruta.write_text(f'''<!doctype html><html lang="{idioma_html}"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{titulo}</title><style>
:root{{--f:#fff;--t:#1d2433;--b:#d9dee8;--a:#1f5fbf}}
@media (prefers-color-scheme:dark){{:root{{--f:#14171d;--t:#e6e9ef;--b:#2c3340;--a:#7fb0ff}}}}
body{{margin:0;padding:16px;font:15px/1.45 system-ui,sans-serif;background:var(--f);color:var(--t)}}
.g{{display:grid;grid-template-columns:repeat(auto-fill,minmax(300px,1fr));gap:12px}}
.t{{border:1px solid var(--b);border-radius:10px;overflow:hidden}} .t img{{width:100%;display:block}}
.d{{padding:8px 10px}} .c{{font-weight:700;color:var(--a)}} small{{opacity:.8}}
table{{border-collapse:collapse;width:100%;margin-bottom:16px}} td,th{{border:1px solid var(--b);padding:6px 8px;text-align:left;vertical-align:top}}
</style></head><body><h1>{encabezado}</h1>
<p>{html.escape(_meta_text(meta, "producto", idioma))}<br>{ind["filas"]} {n_filas} · {ind["duracion_s"]} s · {mano_d}: {pct("MD")} · {mano_i}: {pct("MI")}</p>
<h2>{propuestas_titulo}</h2><ul>{prop}</ul>{err}<h2>{filas_titulo}</h2><div class="g">{"".join(tarjetas)}</div></body></html>''', encoding='utf-8')


if __name__ == '__main__':
    import sys
    pred = json.loads(Path(sys.argv[1]).read_text(encoding='utf-8'))
    filas = unir_filas(pred)
    ind = indicadores(filas)
    print(json.dumps(ind, ensure_ascii=False))
    print(propuestas(ind))
